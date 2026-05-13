"""
Iteration 39 tests:
- Premade courses inline view endpoint (PDF, DOCX, PPTX/etc.)
- /api/advanced/tts JSON body short text path
"""
import os
import io
import base64
import pytest
import requests

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', 'https://knowledge-sync-9.preview.emergentagent.com').rstrip('/')
API = f"{BASE_URL}/api"

TEACHER = {"email": "demo.teacher@eduos.edu", "password": "demo123"}
ORG_ADMIN = {"email": "test.orgadmin@cognify.edu", "password": "admin123"}


def _login(creds):
    r = requests.post(f"{API}/auth/login", json=creds, timeout=20)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    data = r.json()
    token = data.get("access_token") or data.get("token")
    assert token, f"no token in: {data}"
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def teacher_headers():
    return _login(TEACHER)


@pytest.fixture(scope="module")
def admin_headers():
    return _login(ORG_ADMIN)


# ---------- Premade course listing ----------
def test_list_premade_courses(teacher_headers):
    r = requests.get(f"{API}/premade-courses", headers=teacher_headers, timeout=20)
    assert r.status_code == 200, r.text
    data = r.json()
    # Some endpoints return list, others envelope
    if isinstance(data, dict):
        docs = data.get("documents", [])
    else:
        docs = data
    assert isinstance(docs, list)
    print(f"Found {len(docs)} premade courses")


def _get_docs(headers):
    r = requests.get(f"{API}/premade-courses", headers=headers, timeout=20)
    data = r.json()
    docs = data.get("documents", []) if isinstance(data, dict) else data
    return docs


# ---------- Upload helpers (if no existing docs of a type) ----------
def _ensure_pdf(headers):
    docs = _get_docs(headers)
    for d in docs:
        if (d.get("filename") or "").lower().endswith(".pdf"):
            return d["id"]
    # Upload tiny PDF
    pdf_bytes = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
                 b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
                 b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]>>endobj\n"
                 b"xref\n0 4\n0000000000 65535 f \n0000000010 00000 n \n0000000053 00000 n \n0000000099 00000 n \n"
                 b"trailer<</Size 4/Root 1 0 R>>\nstartxref\n149\n%%EOF\n")
    files = {"file": ("TEST_iter39.pdf", pdf_bytes, "application/pdf")}
    data = {"title": "TEST_iter39 pdf", "category": "General", "description": "test"}
    r = requests.post(f"{API}/premade-courses/upload", headers=headers, files=files, data=data, timeout=30)
    assert r.status_code in (200, 201), r.text
    return r.json()["id"]


def _ensure_docx(headers):
    docs = _get_docs(headers)
    for d in docs:
        if (d.get("filename") or "").lower().endswith(".docx"):
            return d["id"]
    # Minimal valid DOCX (zip with required parts) - use python-docx if available, else skip
    try:
        from docx import Document
        bio = io.BytesIO()
        doc = Document()
        doc.add_paragraph("Iter39 docx test")
        doc.save(bio)
        bio.seek(0)
        files = {"file": ("TEST_iter39.docx", bio.getvalue(),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        data = {"title": "TEST_iter39 docx", "category": "General", "description": "test"}
        r = requests.post(f"{API}/premade-courses/upload", headers=headers, files=files, data=data, timeout=30)
        assert r.status_code in (200, 201), r.text
        return r.json()["id"]
    except ImportError:
        pytest.skip("python-docx not installed; cannot synthesize DOCX")


def _ensure_pptx(headers):
    docs = _get_docs(headers)
    for d in docs:
        if (d.get("filename") or "").lower().endswith(".pptx"):
            return d["id"]
    try:
        from pptx import Presentation
        bio = io.BytesIO()
        p = Presentation()
        slide = p.slides.add_slide(p.slide_layouts[0])
        slide.shapes.title.text = "Iter39 pptx test"
        p.save(bio)
        bio.seek(0)
        files = {"file": ("TEST_iter39.pptx", bio.getvalue(),
                          "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
        data = {"title": "TEST_iter39 pptx", "category": "General", "description": "test"}
        r = requests.post(f"{API}/premade-courses/upload", headers=headers, files=files, data=data, timeout=30)
        assert r.status_code in (200, 201), r.text
        return r.json()["id"]
    except ImportError:
        pytest.skip("python-pptx not installed; cannot synthesize PPTX")


# ---------- View endpoint ----------
def test_view_pdf_inline(teacher_headers):
    doc_id = _ensure_pdf(teacher_headers)
    r = requests.get(f"{API}/premade-courses/{doc_id}/view", headers=teacher_headers, timeout=30)
    assert r.status_code == 200, r.text
    assert r.headers.get("content-type", "").startswith("application/pdf"), r.headers
    assert "inline" in (r.headers.get("content-disposition") or "").lower()
    assert r.content[:4] == b"%PDF"


def test_view_docx_inline_original_bytes(teacher_headers):
    doc_id = _ensure_docx(teacher_headers)
    r = requests.get(f"{API}/premade-courses/{doc_id}/view", headers=teacher_headers, timeout=30)
    assert r.status_code == 200, r.text
    ct = r.headers.get("content-type", "")
    assert "wordprocessingml" in ct or "officedocument" in ct, ct
    assert "inline" in (r.headers.get("content-disposition") or "").lower()
    # DOCX is a zip; first 2 bytes are PK
    assert r.content[:2] == b"PK"


def test_view_pptx_converted_to_pdf(teacher_headers):
    doc_id = _ensure_pptx(teacher_headers)
    r = requests.get(f"{API}/premade-courses/{doc_id}/view", headers=teacher_headers, timeout=120)
    assert r.status_code == 200, r.text
    assert r.headers.get("content-type", "").startswith("application/pdf"), r.headers
    assert "inline" in (r.headers.get("content-disposition") or "").lower()
    assert r.content[:4] == b"%PDF"


def test_view_unauthenticated_blocked():
    r = requests.get(f"{API}/premade-courses/some-id/view", timeout=10)
    assert r.status_code in (401, 403)


def test_view_not_found(teacher_headers):
    r = requests.get(f"{API}/premade-courses/nonexistent-id-xyz/view", headers=teacher_headers, timeout=10)
    assert r.status_code == 404


def test_download_endpoint_still_disabled(teacher_headers):
    docs = _get_docs(teacher_headers)
    if not docs:
        pytest.skip("no docs to test")
    doc_id = docs[0]["id"]
    r = requests.get(f"{API}/premade-courses/{doc_id}/download", headers=teacher_headers, timeout=10)
    assert r.status_code == 403, f"download should be disabled, got {r.status_code}"


# ---------- /api/advanced/tts ----------
def test_advanced_tts_short_text(teacher_headers):
    """Short text TTS should accept JSON body. May return 503 if LLM budget exceeded."""
    payload = {"text": "Hello world short test.", "voice": "nova", "speed": 1.0}
    r = requests.post(f"{API}/advanced/tts", headers=teacher_headers, json=payload, timeout=60)
    if r.status_code == 503:
        pytest.skip(f"TTS budget exceeded: {r.text[:200]}")
    assert r.status_code == 200, r.text
    data = r.json()
    assert "audio_url" in data or "audio_data" in data, data
    if "chunks" in data:
        assert isinstance(data["chunks"], int)


def test_advanced_tts_rejects_empty(teacher_headers):
    r = requests.post(f"{API}/advanced/tts", headers=teacher_headers, json={"text": "", "voice": "nova"}, timeout=20)
    assert r.status_code in (400, 422), r.text
