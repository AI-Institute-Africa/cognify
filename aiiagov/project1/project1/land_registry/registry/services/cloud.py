import os
import uuid
import hashlib
import io
import base64
import re
from io import BytesIO
from html import unescape
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, Inches
from django.utils import timezone
from django.db import transaction
from django.core.files.base import ContentFile
from registry.models import CloudFolder, CloudFile, CloudFileVersion, FileAccessLog, CloudShare

class DocxHTMLParser(HTMLParser):
    def __init__(self, document):
        super().__init__()
        self.document = document
        self.current_paragraph = self.document.add_paragraph()
        self.current_style = {'bold': False, 'italic': False, 'underline': False}
        self.list_type = None
        self._table_rows = []
        self._current_row = None
        self._collect_cell = False
        self._cell_text = ''

    def handle_starttag(self, tag, attrs):
        if tag in ('p', 'div'):
            self.current_paragraph = self.document.add_paragraph()
            if self.list_type == 'ul':
                self.current_paragraph.style = 'List Bullet'
            elif self.list_type == 'ol':
                self.current_paragraph.style = 'List Number'
        elif tag == 'br':
            self.current_paragraph.add_run('\n')
        elif tag in ('strong', 'b'):
            self.current_style['bold'] = True
        elif tag in ('em', 'i'):
            self.current_style['italic'] = True
        elif tag == 'u':
            self.current_style['underline'] = True
        elif tag == 'h1':
            self.current_paragraph = self.document.add_paragraph(style='Heading 1')
        elif tag == 'h2':
            self.current_paragraph = self.document.add_paragraph(style='Heading 2')
        elif tag == 'h3':
            self.current_paragraph = self.document.add_paragraph(style='Heading 3')
        elif tag == 'ul':
            self.list_type = 'ul'
        elif tag == 'ol':
            self.list_type = 'ol'
        elif tag == 'li':
            style = 'List Bullet' if self.list_type == 'ul' else 'List Number'
            self.current_paragraph = self.document.add_paragraph(style=style)
        elif tag == 'img':
            # get src attribute
            attrs = dict(attrs)
            src = attrs.get('src', '')
            if src.startswith('data:'):
                # data:[<mediatype>][;base64],<data>
                m = re.match(r'data:(?P<mime>[-\w+/]+);base64,(?P<data>.+)', src)
                if m:
                    try:
                        img_data = base64.b64decode(m.group('data'))
                        img_stream = BytesIO(img_data)
                        # Insert image into current paragraph
                        try:
                            run = self.current_paragraph.add_run()
                            run.add_picture(img_stream)
                        except Exception:
                            # fallback: add picture directly to document
                            self.document.add_picture(img_stream)
                    except Exception:
                        pass
            elif src:
                # try to use remote or local path by adding as image if accessible
                try:
                    # if it's a local /media path, attempt to open
                    if src.startswith('/') or src.startswith('.'):
                        path = src.lstrip('/')
                        if os.path.exists(path):
                            self.current_paragraph.add_run().add_picture(path)
                except Exception:
                    pass
        elif tag == 'table':
            # start collecting table rows
            self._table_rows = []
            self._current_row = None
        elif tag == 'tr':
            self._current_row = []
        elif tag in ('td', 'th'):
            # start collecting cell text
            self._collect_cell = True
            self._cell_text = ''

    def handle_endtag(self, tag):
        if tag in ('strong', 'b'):
            self.current_style['bold'] = False
        elif tag in ('em', 'i'):
            self.current_style['italic'] = False
        elif tag == 'u':
            self.current_style['underline'] = False
        elif tag in ('ul', 'ol'):
            self.list_type = None
        elif tag in ('p', 'div', 'li', 'h1', 'h2', 'h3'):
            self.current_paragraph = self.document.add_paragraph()
        elif tag in ('td', 'th'):
            self._collect_cell = False
            if self._current_row is not None:
                self._current_row.append(self._cell_text)
                self._cell_text = ''
        elif tag == 'tr':
            if self._current_row is not None:
                self._table_rows.append(self._current_row)
                self._current_row = None
        elif tag == 'table':
            # create table in document
            if self._table_rows:
                max_cols = max(len(r) for r in self._table_rows)
                table = self.document.add_table(rows=len(self._table_rows), cols=max_cols)
                for r_idx, row in enumerate(self._table_rows):
                    for c_idx, cell_text in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = cell_text
            self._table_rows = []

    def handle_data(self, data):
        if not data:
            return
        text = unescape(data)
        if self._collect_cell:
            self._cell_text += text
            return

        run = self.current_paragraph.add_run(text)
        run.bold = self.current_style.get('bold', False)
        run.italic = self.current_style.get('italic', False)
        run.underline = self.current_style.get('underline', False)
        # Basic font size/family detection from surrounding tags isn't implemented here,
        # but the `create_document` flow can pass inline-styled HTML which this parser
        # will not fully preserve for complex CSS. We support structure, basic inline
        # emphasis, images and tables which covers most in-browser edits from TinyMCE.

class CloudService:
    @staticmethod
    def get_or_create_folder(name, department, parent=None, project_context=None, user=None):
        """
        Ensures a governance folder exists.
        """
        folder, created = CloudFolder.objects.get_or_create(
            name=name,
            department=department,
            parent=parent,
            defaults={
                'project_context': project_context,
                'created_by': user
            }
        )
        return folder

    @staticmethod
    def create_document(folder, name, content, user, commit_message="Created Word document", is_internal_only=True):
        """
        Creates a new .docx file inside the cloud vault from HTML or plain text content.
        """
        if not name.lower().endswith('.docx'):
            name = f"{name}.docx"

        doc = Document()
        if '<' in content and '>' in content:
            parser = DocxHTMLParser(doc)
            parser.feed(content)
        else:
            for line in content.splitlines():
                doc.add_paragraph(line)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        file_bytes = buffer.getvalue()
        content_file = ContentFile(file_bytes, name=name)

        with transaction.atomic():
            cloud_file, created = CloudFile.objects.get_or_create(
                name=name,
                folder=folder,
                defaults={
                    'owner': user,
                    'is_internal_only': is_internal_only
                }
            )

            last_version = cloud_file.versions.order_by('-version_number').first()
            next_version_num = (last_version.version_number + 1) if last_version else 1

            version = CloudFileVersion.objects.create(
                cloud_file=cloud_file,
                version_number=next_version_num,
                file_content=content_file,
                uploaded_by=user,
                commit_message=commit_message,
                file_size=len(file_bytes)
            )

            CloudService.log_access(user, cloud_file, 'Download' if created else 'View', version=version)

        return cloud_file

    @staticmethod
    def upload_file(folder, uploaded_file, user, commit_message="", is_internal_only=True):
        """
        Uploads a new version of a file within a governance context.
        """
        with transaction.atomic():
            # Find or create the CloudFile object
            cloud_file, created = CloudFile.objects.get_or_create(
                name=uploaded_file.name,
                folder=folder,
                defaults={
                    'owner': user,
                    'is_internal_only': is_internal_only
                }
            )

            # Determine next version number
            last_version = cloud_file.versions.order_by('-version_number').first()
            next_version_num = (last_version.version_number + 1) if last_version else 1

            # Create the file version
            version = CloudFileVersion.objects.create(
                cloud_file=cloud_file,
                version_number=next_version_num,
                file_content=uploaded_file,
                uploaded_by=user,
                commit_message=commit_message,
                file_size=uploaded_file.size
            )

            # Log the action
            CloudService.log_access(user, cloud_file, 'Download' if created else 'View', version=version)
            
            return cloud_file

    @staticmethod
    def log_access(user, cloud_file, action, version=None, request=None):
        """
        Saves an audit trail entry.
        """
        ip = None
        ua = None
        if request:
            x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
            if x_forwarded_for:
                ip = x_forwarded_for.split(',')[0]
            else:
                ip = request.META.get('REMOTE_ADDR')
            ua = request.META.get('HTTP_USER_AGENT')

        FileAccessLog.objects.create(
            user=user,
            cloud_file=cloud_file,
            version=version or cloud_file.latest_version,
            action=action,
            ip_address=ip,
            user_agent=ua
        )

    @staticmethod
    def create_share(cloud_file, shared_by, expiry_days=7, shared_with_user=None, shared_with_email=None, can_download=False, can_edit=False):
        """
        Creates a time-limited sharing link.
        """
        access_token = str(uuid.uuid4())
        expiry_date = timezone.now() + timezone.timedelta(days=expiry_days)
        
        share = CloudShare.objects.create(
            cloud_file=cloud_file,
            shared_by=shared_by,
            shared_with_user=shared_with_user,
            shared_with_email=shared_with_email,
            access_token=access_token,
            expiry_date=expiry_date,
            can_download=can_download,
            can_edit=can_edit
        )
        
        CloudService.log_access(shared_by, cloud_file, 'Share')
        return share

    @staticmethod
    def rollback(cloud_file, version_number, user):
        """
        Promotes an old version to the latest.
        """
        target_version = cloud_file.versions.filter(version_number=version_number).first()
        if not target_version:
            return False
            
        with transaction.atomic():
            next_v = cloud_file.latest_version.version_number + 1
            new_v = CloudFileVersion.objects.create(
                cloud_file=cloud_file,
                version_number=next_v,
                file_content=target_version.file_content,
                uploaded_by=user,
                commit_message=f"Rollback to v{version_number}",
                file_size=target_version.file_size,
                checksum=target_version.checksum
            )
            CloudService.log_access(user, cloud_file, 'Rollback', version=new_v)
            return True
